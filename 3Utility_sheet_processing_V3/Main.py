from Service_address_class import ServiceAddress
from Bill_date_class import BillDate
from Effective_dates_class import EffectiveDates
from Due_date_class import DueDate
from pytesseract import pytesseract
from pdf2image import convert_from_path
from PIL import Image
from datetime import date
from openpyxl import load_workbook
from Amount_class import Amount
import os
import numpy as np
import time

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

def process_with_everything(image, company):
    data = []
    objects = create_objects(image, company)
    for item in objects:
        item.find_data_in_image()
        data.append(item.data)
    return data

def process_just_service_address(image, company):
    service_address = ServiceAddress(image, company)
    service_address.find_data_in_image()
    if company == 2:
        service_address.data = None
    if service_address.data:
        print(f"--- Processing {service_address.data} ---")
    return not service_address.data, service_address.data

def create_objects(image, company):
    return [
        BillDate(image, company),
        Amount(image, company),
        EffectiveDates(image, company),
        DueDate(image, company)
    ]

def add_stamp(image):
    stamp = Image.open("Pictures/Entered_no_bg.png").convert("RGBA")
    base = image.convert("RGBA")
    base.paste(stamp, (650, 1280), mask=stamp)
    return base

def make_and_save_pdf(images):
    rgb_images = [img.convert("RGB") for img in images]
    output_dir = os.path.join("PDFs", "processed_pdfs")
    os.makedirs(output_dir, exist_ok=True)
    today_str = date.today().strftime("%Y-%m-%d")
    pdf_filename = f"utilities_{today_str}.pdf"
    pdf_path = os.path.join(output_dir, pdf_filename)
    first_image, *rest = rgb_images
    first_image.save(pdf_path, save_all=True, append_images=rest)
    print(f"PDF created: {pdf_path}")
    return pdf_path

def parse_month(page):
    month_str = str(page[1]).split("/")[0]
    index_to_search_for_month_name = int(month_str) - 1
    return MONTHS[index_to_search_for_month_name]

def open_sheet(wb, tab_name):
    if tab_name in wb.sheetnames:
        print(f"Opened tab: {tab_name}")
        return wb[tab_name]
    print(f"Tab '{tab_name}' not found in workbook.")
    return None

def find_starting_cell(ws, targets):
    for row in ws.iter_rows(min_col=1, max_col=1):
        cell = row[0]
        if cell.value is not None and cell.value in targets:
            print(f"Found {cell.value} at row {cell.row}")
            return cell
    print("No target found")
    return None

def find_month_column(ws, header_row, month_name):
    for col in range(2, ws.max_column + 1):
        header = ws.cell(row=header_row, column=col).value
        if header and str(header).strip() == month_name:
            print(f"Found {month_name} at column {col} in row {header_row}")
            return col, header_row - 1
    print(f"Month '{month_name}' not found")
    return None

def write_values(ws, values_to_write, target_rows, col):
    for i, r in enumerate(target_rows):
        if i < len(values_to_write):
            ws.cell(row=r, column=col, value=values_to_write[i])

def add_to_excel_sheet(page, image):
    if not page or len(page) < 2:
        return image
    try:
        month_name = parse_month(page)
    except Exception as e:
        print(f"Failed to parse month from {page[1] if len(page) > 1 else 'N/A'}: {e}")
        return image

    file_path = "Utilities Billed to Tenants - Sept25.xlsx"
    wb = load_workbook(file_path)
    ws = open_sheet(wb, page[0])
    if not ws:
        return image

    company_index = page[-1]
    target = "Enbridge" if company_index == 0 else "Logan City" if company_index == 1 else None

    if not target:
        return image

    values_to_write = page[1:]
    cell_to_start = find_starting_cell(ws, target)

    if not cell_to_start:
        return image

    header_row = cell_to_start.row - 1
    month_col, base_cell = find_month_column(ws, header_row, month_name)

    if not month_col:
        return image

    base_target_rows = [2, 6, 5, 4]
    target_rows = [num + base_cell for num in base_target_rows]

    write_values(ws, values_to_write, target_rows, month_col)
    wb.save(file_path)
    print(f"Data written successfully for {target} on tab '{page[0]}'.")
    return add_stamp(image)

def get_utility_company(image):
    avg = tuple(int(x) for x in np.array(image).mean(axis=(0, 1)))
    enbridge_min, enbridge_max = (239, 239, 239), (249, 249, 249)
    logan_min, logan_max = (222, 226, 231), (232, 236, 241)
    if all(enbridge_min[i] <= avg[i] <= enbridge_max[i] for i in range(3)):
        return 0
    elif all(logan_min[i] <= avg[i] <= logan_max[i] for i in range(3)):
        return 1
    return 2

def process_images(images):
    list_of_pages = []
    image_list = []
    for image in images:
        print(image)
        page = []
        company = get_utility_company(image)
        if company != 2:
            is_empty, data = process_just_service_address(image, company)
            if not is_empty:
                page.append(data)
                processed_data = process_with_everything(image, company)
                page.extend(processed_data)
                page.append(company)
                image = add_to_excel_sheet(page, image)
                image_list.append(image)
            else:
                image_list.append(image)
        else:
            image_list.append(image)
        print(page)
        list_of_pages.append(page)
    return list_of_pages, image_list

def main():
    print("Initializing...")

    images = convert_from_path(
        "PDFs/Utilis 10-27-2025.pdf",
        poppler_path=r"C:\poppler\poppler-24.08.0\Library\bin",
    )
    print("Pdf Converted")
    pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    list_of_pages, image_list = process_images(images)
    print(list_of_pages)
    make_and_save_pdf(image_list)

if __name__ == '__main__':
    main()
