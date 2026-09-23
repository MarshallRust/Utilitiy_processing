"""
Everything that touches the actual spreadsheet: matching an OCR'd address
to a real tab name, and writing all of a run's results in one load/save
instead of V3's one load/save per page.
"""

import pandas as pd
from openpyxl import load_workbook

EXCEL_PATH = "Utilities Billed to Tenants - Sept25.xlsx"  # TODO: pull from config instead of hardcoding

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
COMPANY_LABELS = {"enbridge": "Enbridge", "logan": "Logan City"}

# Row offset from the company label cell (e.g. the "Enbridge" row) to where
# each field gets written. Same layout V3 used - still tied to the exact
# spreadsheet template, still worth making more robust later, just carried
# over as-is for now.
ROW_OFFSETS = {"bill_date": 0, "due_date": 2, "effective_dates": 3, "amount": 4}

_tab_lookup = None  # cache: {cleaned_tab_name: original_tab_name}


def _clean(s):
    """Lowercase, strip spaces/punctuation - so '230 W 100 N,' and '230w100n'
    compare equal."""
    if not s:
        return ""
    return "".join(ch for ch in s.lower() if ch.isalnum())


def _load_tabs():
    global _tab_lookup
    if _tab_lookup is None:
        excel_file = pd.ExcelFile(EXCEL_PATH)
        _tab_lookup = {_clean(tab): tab for tab in excel_file.sheet_names}
    return _tab_lookup


def resolve_tab_name(address):
    """Given a raw OCR'd address, return the real Excel tab name it matches,
    or None if it doesn't match any tracked tab."""
    if not address:
        return None

    cleaned_address = _clean(address)
    for cleaned_tab, original_tab in _load_tabs().items():
        if cleaned_tab in cleaned_address:
            return original_tab
    return None


def _parse_month(bill_date):
    """'10/13/2025' -> 'Oct' - the header row uses 3-letter month names."""
    month_num = int(bill_date.split("/")[0])
    return MONTHS[month_num - 1]


def _find_company_row(ws, label):
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value == label:
            return row[0].row
    return None


def _find_month_column(ws, header_row, month_name):
    for col in range(2, ws.max_column + 1):
        header = ws.cell(row=header_row, column=col).value
        if header and str(header).strip() == month_name:
            return col
    return None


def write_all_to_excel(results):
    """Writes every page's results into the workbook in one load and one
    save, instead of V3's load/save per page. Pages that don't fit the
    expected layout (no matching tab, no company row, no column for the
    bill's month) are skipped and reported at the end rather than crashing
    the whole run."""
    if not results:
        print("Nothing to write - no processed pages")
        return

    wb = load_workbook(EXCEL_PATH)
    skipped = []

    for result in results:
        if result.address not in wb.sheetnames:
            skipped.append((result.address, "no matching tab in workbook"))
            continue
        ws = wb[result.address]

        label = COMPANY_LABELS.get(result.company)
        company_row = _find_company_row(ws, label)
        if company_row is None:
            skipped.append((result.address, f"no '{label}' row found"))
            continue

        if not result.bill_date:
            skipped.append((result.address, "no bill_date - can't find the month column"))
            continue

        header_row = company_row - 2
        month_name = _parse_month(result.bill_date)
        month_col = _find_month_column(ws, header_row, month_name)
        if month_col is None:
            skipped.append((result.address, f"no column for {month_name}"))
            continue

        for field_name, offset in ROW_OFFSETS.items():
            value = getattr(result, field_name)
            ws.cell(row=company_row + offset, column=month_col, value=value)

    wb.save(EXCEL_PATH)
    print(f"Saved {len(results) - len(skipped)} of {len(results)} results to {EXCEL_PATH}")
    for address, reason in skipped:
        print(f"  skipped {address}: {reason}")
