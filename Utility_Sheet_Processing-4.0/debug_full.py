"""
Throwaway - runs the real process_page logic over every page (no writing,
no stamping) and prints every field's raw value, plus what's actually in
column 1 of one affected tab, so we can see both issues at once.
"""

from openpyxl import load_workbook

from excel_writer import EXCEL_PATH
from main import process_page
from pdf_io import load_pdf_images

PDF_PATH = "PDFs/Utilis 10-27-2025.pdf"

images = load_pdf_images(PDF_PATH)
print(f"Loaded {len(images)} pages\n")

for i, image in enumerate(images):
    result = process_page(image)
    if result is None:
        continue
    print(f"page {i}: company={result.company!r} address={result.address!r}")
    print(f"  bill_date={result.bill_date!r} due_date={result.due_date!r}")
    print(f"  amount={result.amount!r} effective_dates={result.effective_dates!r}")

print()
print("--- column 1 of the '320 E 400 N' tab (first 15 rows) ---")
wb = load_workbook(EXCEL_PATH)
ws = wb["320 E 400 N"]
for row in range(1, 16):
    value = ws.cell(row=row, column=1).value
    print(f"  row {row}: {value!r}")
